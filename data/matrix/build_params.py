"""Сборка data/matrix/params.csv из Приложения 1 ТЗ и наших правил извлечения.

Источник истины по составу и текстам — source/appendix1-matrix-v1.1.xlsx (лист «МАТРИЦА», организаторы).
Типы, регулярные выражения, пороги и шкалы для «понижения класса» — overrides.csv (ведём сами).
Для параметров без строки в overrides тип выводится по единице измерения.
Внешний код организаторов (``KR-055``, каталог parameter_catalog_132 из датасета) выводится из раздела и номера;
если рядом есть data/samples/parameter-code-map.csv (разбор датасета), коды сверяются с ним.

Запуск из корня репозитория:
    uv run --with openpyxl python data/matrix/build_params.py
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "source" / "appendix1-matrix-v1.1.xlsx"
OVERRIDES = HERE / "overrides.csv"
TARGET = HERE / "params.csv"
CODE_MAP = HERE.parent / "samples" / "parameter-code-map.csv"

COLUMNS = [
    "code", "external_code", "section", "parameter_name", "unit", "source_pd", "source_rd", "source_id",
    "trigger_logic", "review_priority", "sp_reference", "gost_reference", "fz_reference",
    "other_normative", "data_type", "min_value", "max_value", "regex_pattern",
    "semantic_anchors", "enum_values", "is_active",
]
# Активность по умолчанию задаёт api (MATRIX_ACTIVE_PARAMS); здесь все выключены.
NUMBER_UNITS = re.compile(
    r"^(м²|м³|м|мм|мм²|ед\.|шт\.|кВт|м³/сут|м³/ч|Гкал/ч|%|‰|мин|л/с|дни|чел\.|тыс\. руб\.|"
    r"Вт/\(м·С\)|м²·С/Вт|кВт·ч/м²|А \(Ампер\)|мин \(EI\))$"
)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def data_type_for(unit: str) -> str:
    if NUMBER_UNITS.match(unit):
        return "number"
    if unit.startswith("Коорд"):
        return "coordinate"
    if " / " in unit:
        return "string"  # составная единица («Марка / Толщина») — сравнение по частям
    if unit.startswith(("Марка", "Класс", "Кат", "Буква", "Степень")):
        return "enum"
    return "string"  # «—», «Статус» — качественное сравнение


# раздел ПД → префикс кода у организаторов: ПЗ → PZ, ИОС4 → IOS4 и т.д.
SECTION_PREFIX = {
    "ПЗ": "PZ", "СПЗУ": "SPZU", "АР": "AR", "КР": "KR", "ППМ": "PPM", "ПОС": "POS", "ОДИ": "ODI",
    "ПОД": "POD", "ЗУ": "ZU", "ООС": "OOS", "СМ": "SM",
    "ИОС1": "IOS1", "ИОС2": "IOS2", "ИОС3": "IOS3", "ИОС4": "IOS4", "ИОС5": "IOS5",
}


def external_code(code: str, section: str) -> str:
    """M-055 + КР → KR-055."""
    return f"{SECTION_PREFIX[section]}-{code.split('-')[1]}"


def name_pattern(name: str) -> str:
    """Название параметра как шаблон: слова по порядку, между ними любые знаки, «е» = «ё».

    «Строительный объем (Общий)» находит «строительный объём, общий» и «Строительный объем (общий)».
    """
    words = re.findall(r"\w+", name)
    return r"\W+".join(re.escape(w).replace("е", "[её]").replace("Е", "[ЕЁ]") for w in words)


def default_rules(row: dict[str, str]) -> None:
    """Правила извлечения из названия — для параметров, которым своих мы ещё не написали.

    Общий разбор ML (extract/regex.py) ищет строку таблицы по якорю, а если её нет на странице —
    шаблон по тексту. Раньше без своих правил было 122 параметра из 132, и их просто нечем было
    найти. Теперь у каждого есть хотя бы название:
    - число — строка таблицы «название | ед. | значение» или «название …: 12,5»;
    - текст, перечисление, координаты — «название: значение;» до точки с запятой.
    Это минимум, а не проверенное извлечение: на настоящих документах строки называются
    по-разному, и синонимы для ходовых параметров надо дописывать в overrides.csv.
    """
    name = row["parameter_name"]
    pattern = name_pattern(name)
    # Синонимы строк из overrides дополняют название, а не вытесняют его: формулировка
    # самой матрицы — тоже законная подпись строки, и документ, написанный по матрице, должен находиться
    title = re.sub(r"\s+", " ", name).strip()
    anchors = [a for a in row["semantic_anchors"].split("|") if a]
    if title.lower().replace("ё", "е") not in {a.lower().replace("ё", "е") for a in anchors}:
        anchors.append(title)
    row["semantic_anchors"] = "|".join(anchors)
    if row["regex_pattern"]:
        return
    if row["data_type"] == "number":
        row["regex_pattern"] = rf"(?i){pattern}[^\d;]{{0,60}}?(-?\d[\d\s]*(?:[.,]\d+)?)"
    else:
        # после названия — закрывающая скобка или кавычка: «… (дренажи): …», «… АИС "ОСИГ": …»
        row["regex_pattern"] = rf"(?i){pattern}[)»\"”]*(?:\s*[,(][^:;]{{0,40}})?\s*[:—–]\s*([^;]{{1,200}})"


def section_of(raw: str) -> str:
    # «Раздел 5. ИОС1» → «ИОС1»
    m = re.match(r"Раздел\s+\d+\.\s*(.+)", raw)
    return (m.group(1) if m else raw).strip()


def main() -> int:
    ws = load_workbook(SOURCE, read_only=True, data_only=True)["МАТРИЦА"]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(h) for h in rows[0]]
    col = {name: header.index(name) for name in header}
    overrides = {r["code"]: r for r in csv.DictReader(OVERRIDES.open(encoding="utf-8"))}

    out = []
    for raw in rows[1:]:
        code = clean(raw[col["Код параметра"]])
        if not code:
            continue
        unit = clean(raw[col["Ед. изм."]])
        row = {
            "code": code,
            "section": (section := section_of(clean(raw[col["Раздел ПД (ПП РФ № 87)"]]))),
            "external_code": external_code(code, section),
            "parameter_name": clean(raw[col["Контролируемый параметр"]]),
            "unit": "" if unit == "—" else unit,
            "source_pd": clean(raw[col["Источник в ПД"]]),
            "source_rd": clean(raw[col["Источник в РД"]]),
            "source_id": clean(raw[col["Источник в ИД"]]),
            "trigger_logic": clean(raw[col["Логика ИИ-связи (предварительный триггер)"]]),
            "review_priority": clean(raw[col["Приоритет экспертной проверки"]]).split(" ")[0] or "MEDIUM",
            "data_type": data_type_for(unit),
            "is_active": "false",
        }
        for key, value in overrides.pop(code, {}).items():
            if key != "code" and value:
                row[key] = value
        row = {c: row.get(c, "") for c in COLUMNS}
        default_rules(row)
        out.append(row)

    if overrides:
        print(f"В overrides.csv есть коды, которых нет в матрице: {', '.join(overrides)}", file=sys.stderr)
        return 1
    if CODE_MAP.exists():
        known = {r["code"]: r["organizer_code"] for r in csv.DictReader(CODE_MAP.open(encoding="utf-8"))}
        wrong = [r["code"] for r in out if known.get(r["code"], r["external_code"]) != r["external_code"]]
        if wrong:
            print(f"Внешние коды расходятся с {CODE_MAP.name}: {', '.join(wrong)}", file=sys.stderr)
            return 1
    with TARGET.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(out)
    print(f"{TARGET.relative_to(HERE.parent.parent)}: {len(out)} параметров")
    return 0


if __name__ == "__main__":
    sys.exit(main())
